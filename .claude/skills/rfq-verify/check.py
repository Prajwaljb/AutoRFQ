"""Checks built RFQs against spec.md "Done when".

Usage: python .claude/skills/rfq-verify/check.py [--out out] [--live]
Reads out/<email name>/RFQ.xlsx and reply.eml for the four sample emails. Exit code 1 if any check fails.
"""
import argparse
import email
import email.policy
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[3]
REFERENCE = ROOT / "reference" / "email-1 reference output.xlsx"
NO_SOURCE_COLUMNS = {"product_id", "Description", "Lookup Status"}
EXPECTED = {
    "email-1-ids-in-body": ["5106912", "5106935", "5106914", "5106966", "5106891", "5106880"],
    "email-2-ids-in-attachment": ["1376044", "1432040", "233900", "4742155", "5106899", "9677514",
                                  "7099537", "225123"],
    "email-3-mixed-ferguson-and-competitor": ["5106878", "5106907", "4570880", "9677531",
                                              "17783-79-3000", "31-024", "MTX-4409-XQ"],
    "email-4-messy-forward": ["5106926", "8067331", "1844552", "4783795", "5106888", "9677533",
                              "17774-79-3000"],
}
NOT_FOUND = {"225123", "9677533", "17783-79-3000", "31-024", "MTX-4409-XQ", "17774-79-3000"}

failures = []


def check(ok, text):
    print(("  PASS " if ok else "  FAIL ") + text)
    if not ok:
        failures.append(text)


def cell(v):
    return "" if v is None else str(v)


def rows(ws):
    return [tuple(cell(v) for v in r) for r in ws.iter_rows(min_row=2, values_only=True)
            if any(v not in (None, "") for v in r)]


def check_email(name, folder, live):
    print(f"{name}:")
    xlsx = folder / "RFQ.xlsx"
    if not xlsx.exists():
        check(False, f"{xlsx} exists")
        return
    wb = openpyxl.load_workbook(xlsx)
    for sheet in ("RFQ", "Sources", "Notes"):
        check(sheet in wb.sheetnames, f"has a {sheet} sheet")
    if "RFQ" not in wb.sheetnames:
        return
    ws = wb["RFQ"]
    header = [c.value for c in ws[1]]
    rfq = rows(ws)
    ids = [r[0] for r in rfq]
    check(ids == EXPECTED[name], f"items in order, once each: {ids}")

    status = {r[0]: r[header.index("Lookup Status")] for r in rfq}
    expect_nf = NOT_FOUND - ({"225123"} if live else set())
    for pid in ids:
        want = "Not found" if pid in expect_nf else "Found"
        check(status[pid] == want, f"{pid} is {want} (got {status[pid]!r})")

    non_text = [c.coordinate for r in ws.iter_rows(min_row=2) for c in r
                if c.value is not None and not isinstance(c.value, str)]
    check(not non_text, f"every RFQ value is text (not text: {non_text[:5]})")

    if "Sources" in wb.sheetnames:
        sourced = {(r[0], r[1]) for r in rows(wb["Sources"])}
        missing = [(r[0], h) for r in rfq for h, v in zip(header, r)
                   if v and h not in NO_SOURCE_COLUMNS and (r[0], h) not in sourced]
        check(not missing, f"every filled spec cell has a Sources row (missing: {missing[:5]})")

    reply = folder / "reply.eml"
    if reply.exists():
        msg = email.message_from_bytes(reply.read_bytes(), policy=email.policy.default)
        names = [p.get_filename() for p in msg.iter_attachments()]
        check("RFQ.xlsx" in names, f"reply.eml has RFQ.xlsx attached (got {names})")
    else:
        check(False, "reply.eml exists")

    if name == "email-1-ids-in-body" and not live:
        ref = openpyxl.load_workbook(REFERENCE)
        check([c.value for c in ref["RFQ"][1]] == header, "RFQ header equals the reference")
        diffs = [(a[0], h, a[i], b[i]) for a, b in zip(rfq, rows(ref["RFQ"]))
                 for i, h in enumerate(header) if a[i] != b[i]]
        check(not diffs, f"RFQ values equal the reference ({len(diffs)} differ: {diffs[:3]})")
        got, want = rows(wb["Sources"]), rows(ref["Sources"])
        check(sorted(got) == sorted(want),
              f"Sources rows equal the reference ({len(set(got) ^ set(want))} differ)")
        check(got == want, "Sources rows in the reference order")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="out")
    ap.add_argument("--live", action="store_true", help="outputs were built live: 225123 must be Found")
    args = ap.parse_args()
    out = (ROOT / args.out).resolve()
    for name in EXPECTED:
        check_email(name, out / name, args.live)
    print(f"\n{'ALL PASSED' if not failures else f'{len(failures)} FAILED'}")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
