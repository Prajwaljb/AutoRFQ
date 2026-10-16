"""Fills templates/rfq_template.xlsx: RFQ, Sources and Notes sheets (values only; style.py formats)."""
import openpyxl

from rfq import config
from rfq.infra.io import clean_text, save_workbook
from rfq.sources.catalog import KitBroken

TEMPLATE_ONLY_SHEETS = ("How to fill",)  # instructions for people filling it by hand; not in the reference


def load_template(path=None):
    """The template, checked: RFQ and Sources sheets with the expected headers."""
    path = path or config.TEMPLATE
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        raise KitBroken(f"Template {path} is missing or not a valid .xlsx ({type(e).__name__}).") from e
    for sheet, columns in (("RFQ", config.RFQ_COLUMNS), ("Sources", config.SOURCES_COLUMNS)):
        if sheet not in wb.sheetnames:
            raise KitBroken(f"Template {path} has no '{sheet}' sheet.")
        header = [c.value for c in wb[sheet][1]][:len(columns)]
        if header != columns:
            raise KitBroken(f"Template {path}: the '{sheet}' header is not {columns}.")
    return wb


def _write(ws, row, col, value, result=None, item=None, column=""):
    """Writes text. Text starting with '=' stays text: only the Image column is ever a formula."""
    text, cut = clean_text(value)
    if cut and result is not None:
        result.add_note(item, "output", f"{column}: value cut to {config.MAX_CELL_CHARS:,} characters.")
    cell = ws.cell(row=row, column=col, value=text)
    if text.startswith("="):
        cell.data_type = "s"
    return cell


def image_formula(url):
    """=IMAGE("<url>") for an https URL, else None."""
    url = (url or "").strip()
    if not url.lower().startswith("https://") or any(ch in url for ch in "\r\n"):
        return None
    return '=IMAGE("' + url.replace('"', '""') + '")'


def source_rows(item):
    """Sources rows for one item: every filled spec cell, in template column order."""
    rows = []
    if not item.found:
        return rows
    for column in config.RFQ_COLUMNS:
        if column in config.NO_SOURCE_COLUMNS or column not in item.cells:
            continue
        cell = item.cells[column]
        rows.append([item.product_id, column, cell.value, cell.source_type, cell.confidence, cell.source])
    return rows


def build(result, template=None):
    """The filled workbook for a run (not yet saved)."""
    wb = load_template(template)
    for name in TEMPLATE_ONLY_SHEETS:
        if name in wb.sheetnames:
            del wb[name]

    ws = wb["RFQ"]
    col_of = {name: i + 1 for i, name in enumerate(config.RFQ_COLUMNS)}
    for r, item in enumerate(result.items, start=2):
        _write(ws, r, col_of["product_id"], item.product_id)
        if item.found:
            for column in config.SPEC_COLUMNS:
                cell = item.cells.get(column)
                if cell is None or cell.value == "":
                    continue
                if column == "Image":
                    formula = image_formula(cell.value)
                    if formula:
                        ws.cell(row=r, column=col_of["Image"], value=formula)
                    continue
                _write(ws, r, col_of[column], cell.value, result, item, column)
        _write(ws, r, col_of["Lookup Status"], item.lookup_status)

    src = wb["Sources"]
    r = 2
    for item in result.items:
        for row in source_rows(item):
            for c, value in enumerate(row, start=1):
                _write(src, r, c, value)
            r += 1

    notes = wb.create_sheet("Notes")
    for c, name in enumerate(config.NOTES_COLUMNS, start=1):
        notes.cell(row=1, column=c, value=name)
    notes.freeze_panes = "A2"
    for r, note in enumerate(result.sorted_notes(), start=2):
        _write(notes, r, 1, note.product_id)
        _write(notes, r, 2, note.text)
    return wb


def save(wb, path):
    """Saves atomically; returns the path written (a "(2)" copy if RFQ.xlsx is open in Excel)."""
    return save_workbook(wb, path)
