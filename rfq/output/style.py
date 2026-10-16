"""Formatting for RFQ.xlsx in the Ferguson palette. Formatting only: never changes a value."""
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from rfq import config

NAVY, MIST, STEEL, WHITE = "00446A", "E6ECF0", "B2C7D2", "FFFFFF"
FOUND_FILL, FOUND_TEXT = "DCFCE7", "166534"
NOT_FOUND_FILL, NOT_FOUND_TEXT = "FFEDD5", "9A3412"
CONFIDENCE_FILL = {"high": "DCFCE7", "medium": "FEF9C3", "low": "FEE2E2"}
FONT = "Inter"  # brand font; Excel substitutes a sans-serif where it is not installed

HEADER_FONT = Font(name=FONT, bold=True, color=WHITE)
HEADER_FILL = PatternFill("solid", start_color=NAVY, end_color=NAVY)
HEADER_ALIGN = Alignment(vertical="center", wrap_text=True)
GRID = Border(bottom=Side(style="thin", color=STEEL))
URL_RE = re.compile(r"https?://[^\s\"')]+")

RFQ_WIDTHS = {"product_id": 16, "Description": 44, "Product details": 60, "Application": 24, "ASTM": 30,
              "Lettering": 26, "Image": 14, "Picture": 34, "Lookup Status": 13, "Temperature Range": 18,
              "Material": 20, "OSHA": 22}
WRAP_COLUMNS = {"Description", "Product details", "Application", "ASTM", "Lettering", "Material", "OSHA"}


def _header(ws, height=30):
    for cell in ws[1]:
        if cell.value is not None:
            cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, HEADER_ALIGN
    ws.row_dimensions[1].height = height


def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def style_rfq(ws, n_rows):
    _header(ws)
    for i, name in enumerate(config.RFQ_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = RFQ_WIDTHS.get(name, 15)
    status_col = config.RFQ_COLUMNS.index("Lookup Status") + 1
    last = get_column_letter(len(config.RFQ_COLUMNS))
    ws.freeze_panes = "C2"  # header row and the ID + description columns stay visible
    if n_rows:
        ws.auto_filter.ref = f"A1:{last}{n_rows + 1}"
    for r in range(2, n_rows + 2):
        ws.row_dimensions[r].height = 60  # room for the IMAGE() thumbnail
        not_found = ws.cell(row=r, column=status_col).value == config.NOT_FOUND
        for c, name in enumerate(config.RFQ_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=name in WRAP_COLUMNS)
            cell.border = GRID
            if not_found:
                cell.fill = _fill(MIST)
            if name == "product_id":
                cell.font = Font(name="Consolas", bold=True, color=NAVY)
        status = ws.cell(row=r, column=status_col)
        if status.value == config.FOUND:
            status.fill, status.font = _fill(FOUND_FILL), Font(name=FONT, bold=True, color=FOUND_TEXT)
        elif not_found:
            status.fill, status.font = _fill(NOT_FOUND_FILL), Font(name=FONT, bold=True, color=NOT_FOUND_TEXT)


def style_sources(ws):
    _header(ws, 22)
    for col, width in zip("ABCDEF", (12, 18, 40, 22, 12, 80)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:F{ws.max_row}"
    for r in range(2, ws.max_row + 1):
        for c in range(1, 7):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=c in (3, 6))
            ws.cell(row=r, column=c).border = GRID
        conf = ws.cell(row=r, column=5)
        if conf.value in CONFIDENCE_FILL:
            conf.fill = _fill(CONFIDENCE_FILL[conf.value])
        source = ws.cell(row=r, column=6)
        m = URL_RE.search(str(source.value or ""))
        if m:  # clickable; the cell keeps its text
            source.hyperlink = m.group(0)
            source.font = Font(name=FONT, color=NAVY, underline="single")


def style_notes(ws):
    _header(ws, 22)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(name="Consolas", color=NAVY)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")


def apply(wb, n_items):
    style_rfq(wb["RFQ"], n_items)
    style_sources(wb["Sources"])
    if "Notes" in wb.sheetnames:
        style_notes(wb["Notes"])
    return wb
