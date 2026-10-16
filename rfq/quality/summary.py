"""The Summary sheet: status, counts, fill rate, request lines, alerts, flags and AI suggestions.

No timings: the same RFQ always gives the same Summary.
"""
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from rfq import config

FILL_COLUMNS = [c for c in config.SPEC_COLUMNS if c != "Image"]  # Image mirrors Picture
NAVY, MIST = "00446A", "E6ECF0"
TITLE = Font(name="Inter", bold=True, size=14, color=NAVY)
SECTION = Font(name="Inter", bold=True, size=12, color=NAVY)
HEAD_FONT = Font(name="Inter", bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", start_color=NAVY, end_color=NAVY)
STATUS_TEXT = {"ready": "Ready", "needs_review": "Needs human review", "approved": "Approved",
               "sent": "Sent", "failed": "Processing error", "fallback_sent": "Processing error"}


def fill_rates(result):
    """(overall %, [(column, filled, found products, % of found, % of all requested)]) over found items."""
    found = [i for i in result.items if i.found]
    requested = len(result.items)
    rows = []
    filled_total = 0
    for column in FILL_COLUMNS:
        filled = sum(1 for i in found if column in i.cells and i.cells[column].value)
        filled_total += filled
        rows.append((column, filled, len(found), round(100 * filled / len(found), 1) if found else 0.0,
                     round(100 * filled / requested, 1) if requested else 0.0))
    cells = len(found) * len(FILL_COLUMNS)
    return (round(100 * filled_total / cells, 1) if cells else 0.0), rows


def product_fill(result):
    """[(product_id, description, filled, total, %)] for every found product."""
    out = []
    for item in result.items:
        if not item.found:
            continue
        filled = sum(1 for c in FILL_COLUMNS if c in item.cells and item.cells[c].value)
        desc = item.cells["Description"].value if "Description" in item.cells else ""
        out.append((item.product_id, desc, filled, len(FILL_COLUMNS), round(100 * filled / len(FILL_COLUMNS), 1)))
    return out


def spec_cells(result):
    found = [i for i in result.items if i.found]
    filled = sum(1 for i in found for c in FILL_COLUMNS if c in i.cells and i.cells[c].value)
    return filled, len(found) * len(FILL_COLUMNS)


def _money(text):
    try:
        return float(str(text).replace("$", "").replace(",", ""))
    except ValueError:
        return None


def request_lines(result):
    """[(id, description, qty, unit price, line total, price source)] and the estimated total."""
    lines, total = [], 0.0
    for item in result.items:
        if not item.found:
            continue
        price = _money(item.extra.get("Price (USD)", ""))
        qty_num = None
        if item.qty:
            try:
                qty_num = float(item.qty.split()[0])
            except ValueError:
                qty_num = None
        line_total = round(price * qty_num, 2) if price is not None and qty_num is not None else None
        if line_total is not None:
            total += line_total
        source = "ferguson.com (live)" if item.found_via == "live" else f"catalog list price, {config.CATALOG_DATE}"
        desc = item.cells["Description"].value if "Description" in item.cells else ""
        lines.append((item.product_id, desc, item.qty, price, line_total, source if price is not None else ""))
    return lines, round(total, 2)


def alerts(result):
    out = []
    for item in result.items:
        if item.found and item.extra.get("Availability") == "OutOfStock":
            out.append((item.product_id, "Out of stock: lead-time risk."))
    for item in result.items:
        if item.found and item.found_via != "live" and item.extra.get("Price (USD)"):
            out.append((item.product_id, f"Price is the catalog list price as of {config.CATALOG_DATE}, not live."))
    for item in result.items:
        if not item.found:
            out.append((item.product_id, item.reason or "Not found."))
    return out


def status_line(result):
    text = STATUS_TEXT.get(result.status.value, result.status.value)
    if result.reviewed_by:
        text = f"Reviewed by {result.reviewed_by}, {result.reviewed_at}"
    return text


def _table(ws, row, headers, rows):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    for r in rows:
        row += 1
        for c, v in enumerate(r, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
    return row + 2


def _section(ws, row, title):
    ws.cell(row=row, column=1, value=title).font = SECTION
    return row + 1


RULE_GROUP = {"missing": "missing", "missing_key_value": "missing", "wrong_unit": "unit", "not_numeric": "unit",
              "out_of_range": "suspicious", "description_mismatch": "suspicious", "placeholder": "suspicious",
              "application_mismatch": "suspicious", "not_found": "not found"}


def add_sheet(wb, result):
    """The RFQ quality summary, as on the deck: overview, fill rate per column and per product, every flag."""
    ws = wb.create_sheet("Summary", index=1 if "RFQ" in wb.sheetnames else None)
    for col, width in zip("ABCDEFGH", (34, 44, 18, 16, 34, 18, 14, 12)):
        ws.column_dimensions[col].width = width
    ws["A1"] = "RFQ quality summary"
    ws["A1"].font = TITLE
    counts = result.counts()
    overall, by_column = fill_rates(result)
    filled, total_cells = spec_cells(result)
    groups = {"missing": 0, "unit": 0, "suspicious": 0, "not found": 0}
    for f in result.flags:
        if not f.resolved:
            groups[RULE_GROUP.get(f.rule, "suspicious")] += 1

    row = _section(ws, 3, "Overview")
    overview = [
        ("Status", status_line(result)),
        ("Products requested", counts["items"]),
        ("Found", counts["found"]),
        ("Found (web, live ferguson.com)", counts["by_source"]["live"]),
        ("Not found", counts["not_found"]),
        ("Spec cells filled (found products)", f"{filled} of {total_cells}"),
        ("Overall fill rate (found products)", f"{overall}%"),
        ("Flagged: missing (by code)", groups["missing"]),
        ("Flagged: unit (by code)", groups["unit"]),
        ("Flagged: suspicious (by code)", groups["suspicious"]),
        ("Flagged: suspicious (by model)", len(result.ai_suggestions)),
        ("Flagged: not found", groups["not found"]),
        ("Units converted", result.units_converted),
        ("Suspicious-value check (LLM)", result.ai_status or "not run"),
    ]
    row = _table(ws, row, ["Item", "Value"], overview)
    status_cell = ws.cell(row=5, column=2)
    status_cell.font = Font(name="Inter", bold=True, color="9A3412" if result.status.value == "needs_review"
                            else "166534")

    row = _section(ws, row, "Every flag (code = rule, model = LLM; the value is never changed)")
    flag_rows = [(f.product_id, f.column, f.value, f.issue, f.rule, f.severity, "code",
                  "resolved" if f.resolved else "open")
                 for f in result.flags if f.severity != "low" and f.rule != "not_found"]
    flag_rows += [(f.product_id, f.column, f.value, f.issue, "llm_check", f.severity, "model", "advisory")
                  for f in result.ai_suggestions]
    missing = sorted({(f.product_id, f.column) for f in result.flags if f.rule == "missing"})
    if missing:
        by_item = {}
        for pid, column in missing:
            by_item.setdefault(pid, []).append(column)
        flag_rows += [(pid, ", ".join(cols), "", f"{len(cols)} blank cell{'s' if len(cols) != 1 else ''}",
                       "missing", "low", "code", "open") for pid, cols in by_item.items()]
    row = _table(ws, row, ["Item", "Column", "Value", "Issue", "Rule", "Severity", "By", "State"],
                 flag_rows or [("", "", "", "No flags", "", "", "", "")])

    row = _section(ws, row, "Fill rate per column")
    first = row
    row = _table(ws, row, ["Column", "Filled", "Found products", "Fill rate (found)", "Fill rate (all requested)"],
                 [(c, f, t, f"{p}%", f"{a}%") for c, f, t, p, a in by_column])
    if by_column and any(t for _, _, t, _, _ in by_column):
        for r, (_, _, _, p, _) in enumerate(by_column, start=first + 1):
            ws.cell(row=r, column=9, value=p)  # numbers for the chart (column I)
        chart = BarChart()
        chart.type = "bar"
        chart.title = "Fill rate by column (%)"
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 100
        chart.legend = None
        data = Reference(ws, min_col=9, min_row=first + 1, max_row=first + len(by_column))
        cats = Reference(ws, min_col=1, min_row=first + 1, max_row=first + len(by_column))
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(cats)
        chart.height, chart.width = 12, 16
        chart.series[0].graphicalProperties.solidFill = NAVY
        ws.add_chart(chart, f"G{first}")

    row = _section(ws, row, "Fill rate per product")
    row = _table(ws, row, ["Item", "Description", "Filled", "Spec cells", "Fill rate"],
                 [(pid, d, f, t, f"{p}%") for pid, d, f, t, p in product_fill(result)] or [("", "No found products", "", "", "")])

    lines, total = request_lines(result)
    row = _section(ws, row, "Request lines")
    row = _table(ws, row, ["Item", "Description", "Qty", "Unit price (USD)", "Price source", "Line total (USD)"],
                 [(l[0], l[1], l[2], l[3], l[5], l[4]) for l in lines])
    ws.cell(row=row - 1, column=5, value="Estimated total (USD)").font = Font(name="Inter", bold=True)
    ws.cell(row=row - 1, column=6, value=total).font = Font(name="Inter", bold=True)
    row += 1

    row = _section(ws, row, "Alerts")
    row = _table(ws, row, ["Item", "Alert"], alerts(result) or [("", "None")])

    alt_rows = []
    for item in result.items:
        if not item.found:
            for o in item.extra.get("alternatives") or []:
                alt_rows.append((item.product_id, o["product_id"], o["description"], ", ".join(o["matched"]),
                                 _money(o.get("price", "")), o.get("availability", "")))
    if alt_rows:
        row = _section(ws, row, "Suggested Ferguson alternatives (not on the RFQ until the customer confirms)")
        row = _table(ws, row, ["Requested", "Ferguson item", "Description", "Matches", "Unit price (USD)",
                               "Availability"], alt_rows)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    return ws
