import email
import email.policy
import subprocess
import sys
from email.message import EmailMessage
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
EMAILS = ROOT / "emails"
sys.path.insert(0, str(ROOT))


def run_build(eml, out, *args, env=None):
    """Runs rfq/build.py the way mailbot.py does. Returns the CompletedProcess."""
    return subprocess.run([sys.executable, str(ROOT / "rfq" / "build.py"), str(eml), "--out", str(out), *args],
                          cwd=ROOT, capture_output=True, text=True, encoding="utf-8", env=env, timeout=300)


def sheet_rows(path, sheet):
    """Non-empty rows of a sheet as tuples of strings ('' for empty)."""
    wb = openpyxl.load_workbook(path)
    rows = []
    for row in wb[sheet].iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in row):
            rows.append(tuple("" if v is None else str(v) for v in row))
    return rows


def all_values(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: [tuple(c.value for c in row) for row in ws.iter_rows()] for ws in wb.worksheets}


def make_email(path, body="", subject="RFQ test", html=None, attachments=(), charset=None, raw_body=None):
    """Writes a small .eml fixture (no real customer data)."""
    msg = EmailMessage(policy=email.policy.default)
    msg["From"] = "Test Buyer <buyer@example.com>"
    msg["To"] = "rfq@example.com"
    msg["Subject"] = subject
    msg["Message-ID"] = "<fixture@example.com>"
    if html is not None and not body:
        msg.set_content(html, subtype="html")
    else:
        msg.set_content(body or "Hello")
        if html is not None:
            msg.add_alternative(html, subtype="html")
    for name, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    raw = msg.as_bytes()
    if raw_body is not None:
        raw = raw_body
    path.write_bytes(raw)
    return path


@pytest.fixture(scope="session")
def offline_builds(tmp_path_factory):
    """The four sample emails built once, offline."""
    out = tmp_path_factory.mktemp("offline")
    for eml in sorted(EMAILS.glob("*.eml")):
        proc = run_build(eml, out, "--offline")
        assert proc.returncode == 0, proc.stderr
    return out
