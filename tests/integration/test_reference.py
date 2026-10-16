"""spec.md "Done when": email 1 equals the reference; unknown items are Not found; nothing is dropped."""
import email
import email.policy
import json

import openpyxl
import pytest

from conftest import ROOT, run_build, sheet_rows

REFERENCE = ROOT / "reference" / "email-1 reference output.xlsx"
NOT_FOUND = {"225123", "9677533", "17783-79-3000", "31-024", "MTX-4409-XQ", "17774-79-3000"}


def test_email1_rfq_equals_reference(offline_builds):
    ours = offline_builds / "email-1-ids-in-body" / "RFQ.xlsx"
    assert sheet_rows(ours, "RFQ") == sheet_rows(REFERENCE, "RFQ")
    ref_ws, our_ws = openpyxl.load_workbook(REFERENCE)["RFQ"], openpyxl.load_workbook(ours)["RFQ"]
    assert [c.value for c in our_ws[1]] == [c.value for c in ref_ws[1]]
    for row in our_ws.iter_rows(min_row=2):
        for cell in row:
            if cell.value is not None:
                assert isinstance(cell.value, str)
                assert (cell.data_type == "f") == (cell.column_letter == "U")  # only Image is a formula


def test_email1_sources_equal_reference_in_order(offline_builds):
    ours = offline_builds / "email-1-ids-in-body" / "RFQ.xlsx"
    assert sheet_rows(ours, "Sources") == sheet_rows(REFERENCE, "Sources")
    assert len(sheet_rows(ours, "Sources")) == 105


def test_email1_sheets_and_status(offline_builds):
    folder = offline_builds / "email-1-ids-in-body"
    assert openpyxl.load_workbook(folder / "RFQ.xlsx").sheetnames[:3] == ["RFQ", "Sources", "Notes"]
    assert sheet_rows(folder / "RFQ.xlsx", "Notes") == []
    assert json.loads((folder / "run.json").read_text(encoding="utf-8"))["status"] == "ready"


def test_email1_without_quality_is_the_same(offline_builds, tmp_path):
    proc = run_build(ROOT / "emails" / "email-1-ids-in-body.eml", tmp_path, "--offline", "--no-quality")
    assert proc.returncode == 0, proc.stderr
    for sheet in ("RFQ", "Sources"):
        assert sheet_rows(tmp_path / "email-1-ids-in-body" / "RFQ.xlsx", sheet) == \
            sheet_rows(offline_builds / "email-1-ids-in-body" / "RFQ.xlsx", sheet)


@pytest.mark.parametrize("name", ["email-2-ids-in-attachment", "email-3-mixed-ferguson-and-competitor",
                                  "email-4-messy-forward"])
def test_not_found_items(offline_builds, name):
    rows = sheet_rows(offline_builds / name / "RFQ.xlsx", "RFQ")
    for row in rows:
        want = "Not found" if row[0] in NOT_FOUND else "Found"
        assert row[-1] == want, row[0]
        if want == "Not found":
            assert all(v == "" for v in row[1:-1]), "a Not found row has only product_id and status"
    ids = [r[0] for r in rows]
    assert len(ids) == len(set(ids))


def test_every_filled_cell_has_a_source(offline_builds):
    for folder in offline_builds.iterdir():
        xlsx = folder / "RFQ.xlsx"
        if not xlsx.exists():
            continue
        header = [c.value for c in openpyxl.load_workbook(xlsx)["RFQ"][1]]
        sourced = {(r[0], r[1]) for r in sheet_rows(xlsx, "Sources")}
        for row in sheet_rows(xlsx, "RFQ"):
            for column, value in zip(header, row):
                if value and column not in ("product_id", "Description", "Lookup Status"):
                    assert (row[0], column) in sourced


def test_reply_eml(offline_builds):
    msg = email.message_from_bytes((offline_builds / "email-4-messy-forward" / "reply.eml").read_bytes(),
                                   policy=email.policy.default)
    assert [p.get_filename() for p in msg.iter_attachments()] == ["RFQ.xlsx"]
    text = msg.get_body(("plain",)).get_content()
    assert "Products found: 5 of 7. Not found: 2." in text
    assert "Closest Ferguson match: 9677539 (3 in. x 1000 ft. Marking Tape in Red" in text
    assert 'Reply "use 5106914"' in text  # the Mutual electric detectable tape -> red electric detector tape
    for word in ("flag", "severity", "DRAFT", " AI ", "fill rate"):
        assert word.lower() not in text.lower()
    assert msg["In-Reply-To"] == "<rfq-sample-4@sourcing.example.com>"
