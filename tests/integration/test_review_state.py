"""Hold -> review -> approve: decisions, undo, stale writes, and the rebuilt RFQ without review marks."""
import email
import email.policy
import json

import openpyxl
import pytest

from conftest import EMAILS, sheet_rows
from rfq import pipeline
from rfq.review import state
from rfq.sources.catalog import CatalogSource


@pytest.fixture
def held(tmp_path, monkeypatch):
    """Email 1 with a wrong width on 5106912: held for review."""
    original = CatalogSource.lookup

    def wrong_width(self, product_id):
        cells = original(self, product_id)
        if product_id == "5106912":
            cells["Width"].value = "60 in"
        return cells
    monkeypatch.setattr(CatalogSource, "lookup", wrong_width)
    result, folder = pipeline.run(EMAILS / "email-1-ids-in-body.eml", tmp_path, mode="offline")
    monkeypatch.setattr(CatalogSource, "lookup", original)
    assert result.status.value == "needs_review"
    return folder


def test_held_rfq_has_review_marks_and_a_clean_reply(held):
    ws = openpyxl.load_workbook(held / "RFQ.xlsx")["RFQ"]
    width = ws.cell(row=2, column=15)
    assert width.value == "60 in" and width.comment is not None
    msg = email.message_from_bytes((held / "reply.eml").read_bytes(), policy=email.policy.default)
    text = msg.get_body(("plain",)).get_content().lower()
    assert "flag" not in text and "review" not in text


def test_decisions_undo_and_stale(held):
    view = state.view(held)
    high = [f for f in view["flags"] if f["severity"] == "high"]
    assert high and not view["can_approve"] and "Resolve" in view["approve_blocked_reason"]
    state.decide(held, high[0]["id"], "keep", "Asha", expected_version=0)
    with pytest.raises(state.Stale):
        state.decide(held, high[0]["id"], "edit", "Ben", value="6 in", expected_version=0)
    state.undo(held, expected_version=1)
    assert state.view(held)["resolved"] == 0
    with pytest.raises(state.ReviewError):
        state.decide(held, high[0]["id"], "edit", "Asha", value="  ")


def test_approve_rebuilds_without_marks(held):
    for f in state.view(held)["flags"]:
        if f["severity"] == "high":
            state.decide(held, f["id"], "edit", "Asha", value="6 in")
    result = state.finalise(held, "Asha")
    assert result.status.value == "approved"
    wb = openpyxl.load_workbook(held / "RFQ.xlsx")
    width = wb["RFQ"].cell(row=2, column=15)
    assert width.value == "6 in" and width.comment is None
    sources = [r for r in sheet_rows(held / "RFQ.xlsx", "Sources") if r[:2] == ("5106912", "Width")]
    assert sources[0][3:5] == ("human_review", "high") and "Asha" in sources[0][5]
    summary = [r for r in sheet_rows(held / "RFQ.xlsx", "Summary") if r[0] == "Status"]
    assert summary[0][1].startswith("Reviewed by Asha")
    run = json.loads((held / "run.json").read_text(encoding="utf-8"))
    assert run["status"] == "approved"
    msg = email.message_from_bytes((held / "reply.eml").read_bytes(), policy=email.policy.default)
    assert "Quality check: checked by our team." in msg.get_body(("plain",)).get_content()
    with pytest.raises(state.ReviewError):
        state.finalise(held, "Asha")  # already approved


def test_substitute_a_not_found_item(tmp_path):
    result, folder = pipeline.run(EMAILS / "email-4-messy-forward.eml", tmp_path, mode="offline")
    with pytest.raises(state.ReviewError):
        state.substitute(folder, "9677533", "5106926", "Asha")  # already on the RFQ
    state.substitute(folder, "9677533", "9677539", "Asha")
    rows = {r[0]: r for r in sheet_rows(folder / "RFQ.xlsx", "RFQ")}
    assert "9677533" not in rows and rows["9677539"][-1] == "Found"
    notes = [r[1] for r in sheet_rows(folder / "RFQ.xlsx", "Notes") if r[0] == "9677539"]
    assert any("Requested 9677533; replaced by Ferguson item 9677539 by Asha" in n for n in notes)
    msg = email.message_from_bytes((folder / "reply.eml").read_bytes(), policy=email.policy.default)
    assert "9677533 (9677539); please confirm" in msg.get_body(("plain",)).get_content()
